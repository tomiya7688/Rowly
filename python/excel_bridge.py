#!/usr/bin/env python3
import json
import sys

from openpyxl import Workbook, load_workbook


def read_request():
    return json.load(sys.stdin)


def stringify(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def export_workbook(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = request["sheet_name"]

    for row_index, row in enumerate(request["rows"], start=1):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.value = value
            cell.data_type = "s"
            cell.number_format = "@"

    workbook.save(request["output_path"])
    json.dump({"ok": True}, sys.stdout)


def import_workbook(request):
    workbook = load_workbook(request["input_path"], data_only=False, read_only=True)
    sheet_name = request.get("sheet_name")
    if sheet_name is None:
        sheet = workbook.active
    else:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"worksheet not found: {sheet_name}")
        sheet = workbook[sheet_name]

    rows = []
    for raw_row in sheet.iter_rows(values_only=True):
        values = list(raw_row)
        while values and values[-1] is None:
            values.pop()
        rows.append([stringify(value) for value in values])

    while rows and not rows[-1]:
        rows.pop()

    json.dump({"rows": rows}, sys.stdout, ensure_ascii=False)


def main():
    if len(sys.argv) != 2 or sys.argv[1] not in {"export", "import"}:
        raise SystemExit("usage: excel_bridge.py <export|import>")

    request = read_request()
    if sys.argv[1] == "export":
        export_workbook(request)
    else:
        import_workbook(request)


if __name__ == "__main__":
    main()
